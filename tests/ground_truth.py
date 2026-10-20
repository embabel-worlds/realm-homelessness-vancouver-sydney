"""
Reconcile reference/ against the sources, read a second way, and write tests/expected.json —
the exact figures tests/views.smoke.mjs then demands of the LIVE views.

The build reads each source once; this reads them again by a different route (a regex over the
City of Sydney page instead of an HTML parser, the curated Vancouver tables summed across rather
than copied, the ABS cubes by cell) and fails on any disagreement. It also checks the internal
arithmetic the sources promise: sheltered + unsheltered = total, the sub-regions add up to the
region, the NSW LGAs add up to the state from 2024.

Run from the realm root:  python3 tests/ground_truth.py
"""
import html as html_lib
import json
import re
import sys
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "source"
problems = []


def check(ok, what):
    if not ok:
        problems.append(what)


figures = {r["data"]["figureId"]: r["data"] for r in yaml.safe_load((ROOT / "reference/homelessness-04-figures.yml").read_text())}
places = {r["data"]["placeId"]: r["data"] for r in yaml.safe_load((ROOT / "reference/homelessness-02-places.yml").read_text())}


def val(fid):
    f = figures.get(fid)
    check(f is not None, f"missing figure {fid}")
    return f and f["value"]


# --- Vancouver: the curated tables must be internally consistent and land in reference/ intact.
c = yaml.safe_load((ROOT / "data/curated/hsabc-greater-vancouver.yml").read_text())
for year, r in c["regional"].items():
    check(r["sheltered"] + r["unsheltered"] == r["total"], f"HSABC {year}: sheltered + unsheltered != total")
    check(val(f"hsabc-pit:van-metro:{year}-03:total") == r["total"], f"HSABC {year} total not in reference")
for i, year in enumerate(c["subregionYears"]):
    s = sum(v[i] or 0 for v in c["subregionTotals"].values())
    # Table 2 itself falls 12-19 short of the regional total in 2005-2014, the years Delta and
    # White Rock are n/a; from 2017 it adds up exactly. See the curated file's header.
    want = c["regional"][year]["total"]
    check(s == want if year >= 2017 else want - 20 < s < want, f"HSABC {year}: sub-regions sum to {s}, region is {want}")
for sub, r in c["subregions2025"].items():
    check(r["ewr"] + r["shelters"] + r["nfa"] == r["sheltered"], f"HSABC 2025 {sub}: sheltered split does not add up")
    check(r["sheltered"] + r["unsheltered"] == r["total"], f"HSABC 2025 {sub}: sheltered + unsheltered != total")
check(sum(r["unsheltered"] for r in c["subregions2025"].values()) == c["regional"][2025]["unsheltered"], "HSABC 2025 unsheltered by sub-region != regional")
check(sum(c["subregionUnsheltered2023"].values()) == c["regional"][2023]["unsheltered"], "HSABC 2023 unsheltered by sub-region != regional")

# --- City of Sydney: regex over the page, independent of the build's HTML parser.
html = (SOURCE / "cos_streetcounts.html").read_text()
table = html[html.index("<table"):html.index("</table>")]
text = " ".join(html_lib.unescape(re.sub(r"<[^>]+>", " ", table)).split())
cos = re.findall(r"(February|August) (\d{4}) ([\d,]+) ([\d,]+)", text)
check(len(cos) >= 27, f"City of Sydney: only {len(cos)} counts in the page's table text")
for month, year, rough, beds in cos:
    period = f"{year}-{'02' if month == 'February' else '08'}"
    check(val(f"cos-street-count:syd-city:{period}:roughSleeping") == int(rough.replace(",", "")), f"CoS {period} rough sleeping")
    check(val(f"cos-street-count:syd-city:{period}:occupiedCrisisBeds") == int(beds.replace(",", "")), f"CoS {period} beds")

# --- NSW: LGAs must add to the state total from 2024 (the paper lists only LGAs still counted).
for year in (2024, 2025, 2026):
    lgas = sum(f["value"] for f in figures.values() if f["sourceId"] == "nsw-street-count" and f["year"] == year and f["placeId"] != "nsw")
    lgas += val(f"cos-street-count:syd-city:{year}-02:roughSleeping")
    check(lgas == val(f"nsw-street-count:nsw:{year}:roughSleeping"), f"NSW {year}: LGAs {lgas} != state total")

# --- ABS: cells read directly.
wb = openpyxl.load_workbook(SOURCE / "abs_20490do007_2021.xlsx", read_only=True, data_only=True)
gs = [r for r in wb["Table_7.7"].iter_rows(values_only=True) if r[0] == "Greater Sydney"]
check(gs[-1][7] == 25141 and val("abs-census:syd-metro:2021-08:total") == 25141, "ABS Greater Sydney 2021 total")
check(gs[0][7] == val("abs-census:syd-metro:2021-08:indigenous"), "ABS Greater Sydney 2021 Indigenous")
wb = openpyxl.load_workbook(SOURCE / "abs_20490do006_2021.xlsx", read_only=True, data_only=True)
for sheet, label, year in (("Table_6.1", "Sydney", 2021), ("Table_6.2", "Sydney (C)", 2016)):
    row = next(r for r in wb[sheet].iter_rows(values_only=True) if r[2] == label)
    check(row[3] == val(f"abs-census:syd-city:{year}-08:improvisedOrSleepingOut"), f"ABS City of Sydney {year} sleeping out")
    check(row[9] == val(f"abs-census:syd-city:{year}-08:total"), f"ABS City of Sydney {year} total")

# --- every relation target exists.
for f in figures.values():
    check(f["placeId"] in places, f"{f['figureId']}: unknown place")
for p in places.values():
    check(not p.get("parent") or p["parent"] in places, f"{p['placeId']}: unknown parent")

if problems:
    print("\n".join("FAIL  " + p for p in problems))
    sys.exit(1)


# --- the exact rows the live views must return.
def latest(place, source, compare):
    fs = [f for f in figures.values() if f["placeId"] == place and f["sourceId"] == source and f.get("compareAs") == compare]
    return max(fs, key=lambda f: f["period"]) if fs else None


expected = {
    "figures": len(figures),
    "places": len(places),
    "sources": 6,
    "latestComparison": sorted(
        [{"compareAs": f["compareAs"], "place": places[f["placeId"]]["name"], "source": f["sourceId"],
          "period": f["period"], "people": f["value"], "per10k": f.get("per10k")}
         for f in (latest(p, s, k) for p in places if places[p]["scale"] in ("city", "metro")
                   for s in {x["sourceId"] for x in figures.values()} for k in ("sleepingOutside", "inCrisisAccommodation", "oneNightTotal"))
         if f], key=lambda r: (r["compareAs"], r["place"], r["source"])),
    "vancouverSubregions2025": {places[f["placeId"]]["name"]: f["value"] for f in figures.values()
                                if f["sourceId"] == "hsabc-pit" and f["year"] == 2025 and f["measure"] == "total"
                                and places[f["placeId"]]["scale"] in ("city", "subregion")},
    "sydneyStreetCountPeriods": len(cos),
    "greaterSydneyStreetCount": {
        str(y): val(f"cos-street-count:syd-city:{y}-02:roughSleeping") + sum(
            f["value"] for f in figures.values() if f["sourceId"] == "nsw-street-count" and f["year"] == y
            and places[f["placeId"]].get("inGreaterSydney")) for y in range(2020, 2027)},
}
(ROOT / "tests/expected.json").write_text(json.dumps(expected, indent=1) + "\n")
print(f"all reconciled; wrote tests/expected.json ({len(expected['latestComparison'])} headline rows)")
